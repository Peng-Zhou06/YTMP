
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import re

def parse_markdown_to_docx(md_content):
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(10.5)
    
    # 设置标题样式
    heading_styles = {
        1: {'size': Pt(16), 'bold': True, 'space_after': Pt(12)},
        2: {'size': Pt(14), 'bold': True, 'space_after': Pt(10)},
        3: {'size': Pt(12), 'bold': True, 'space_after': Pt(8)},
        4: {'size': Pt(11), 'bold': True, 'space_after': Pt(6)},
    }
    
    lines = md_content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # 标题处理
        if line.startswith('#'):
            level = line.count('#')
            if level in heading_styles:
                title_text = line.lstrip('#').strip()
                paragraph = doc.add_paragraph(title_text)
                style = heading_styles[level]
                paragraph.style = f'Heading {level}'
                run = paragraph.runs[0]
                run.font.size = style['size']
                run.bold = style['bold']
                paragraph.space_after = style['space_after']
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # 表格处理
        elif line.startswith('|') and '---' in lines[i+1] if i+1 < len(lines) else False:
            # 解析表格
            table_data = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [cell.strip() for cell in lines[i].split('|')[1:-1]]
                table_data.append(row)
                i += 1
            
            if len(table_data) > 1:
                # 创建表格
                rows = len(table_data)
                cols = len(table_data[0])
                table = doc.add_table(rows=rows, cols=cols)
                table.style = 'Table Grid'
                
                for r, row_data in enumerate(table_data):
                    for c, cell_data in enumerate(row_data):
                        cell = table.cell(r, c)
                        cell.text = cell_data
                        if r == 0:
                            # 表头样式
                            paragraph = cell.paragraphs[0]
                            run = paragraph.runs[0]
                            run.bold = True
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                i -= 1
        
        # 列表处理
        elif line.startswith('- ') or line.startswith('* '):
            # 无序列表
            text = line[2:].strip()
            paragraph = doc.add_paragraph(text)
            paragraph.style = 'List Bullet'
        
        elif re.match(r'^\d+\.', line):
            # 有序列表
            text = re.sub(r'^\d+\.\s*', '', line)
            paragraph = doc.add_paragraph(text)
            paragraph.style = 'List Number'
        
        # 代码块处理
        elif line.startswith('```'):
            # 找到代码块结束
            code_content = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_content.append(lines[i])
                i += 1
            
            if code_content:
                paragraph = doc.add_paragraph('\n'.join(code_content))
                run = paragraph.runs[0]
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
        
        # 分隔线
        elif line.startswith('---'):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(12)
        
        # 普通段落
        elif line.strip():
            paragraph = doc.add_paragraph(line)
            paragraph.space_after = Pt(6)
        
        i += 1
    
    return doc

# 读取Markdown文件
with open('E:/xiaozhou/YTMP/项目设计文档.md', 'r', encoding='utf-8', errors='ignore') as f:
    md_content = f.read()

# 转换为DOCX
doc = parse_markdown_to_docx(md_content)

# 保存DOCX文件
doc.save('E:/xiaozhou/YTMP/YTMP_项目设计文档_完整版.docx')
print("DOCX文档已成功生成！")